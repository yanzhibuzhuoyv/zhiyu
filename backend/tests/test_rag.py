from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.rag import RagService


def test_chinese_retrieval_is_persistent_and_does_not_return_irrelevant_chunks(tmp_path):
    service = RagService(tmp_path)
    document = service.ingest("alice", "笔记.md", "重庆火锅使用花椒和辣椒。".encode())
    service.ingest("alice", "旅行.txt", "杭州西湖适合散步。".encode())

    reopened = RagService(tmp_path)
    results = reopened.search("alice", "火锅 花椒")

    assert reopened.mode == "bm25"
    assert results[0]["document_id"] == document["id"]
    assert results[0]["name"] == "笔记.md"
    assert results[0]["score"] > 0
    assert "火锅" in results[0]["text"]
    assert reopened.search("alice", "天文学") == []
    assert set(document) == {"id", "name", "size", "chunks", "created_at"}
    assert document["size"] == len("重庆火锅使用花椒和辣椒。".encode())


def test_owners_cannot_read_or_delete_each_others_documents(tmp_path):
    service = RagService(tmp_path)
    alice_doc = service.ingest("alice", "secret.txt", b"private astronomy note")
    bob_doc = service.ingest("bob", "other.txt", b"private bicycle note")

    assert [item["id"] for item in service.list_documents("alice")] == [alice_doc["id"]]
    assert [item["id"] for item in service.list_documents("bob")] == [bob_doc["id"]]
    assert service.search("bob", "astronomy") == []
    assert not service.delete("bob", alice_doc["id"])
    assert service.delete("alice", alice_doc["id"])
    assert not service.delete("alice", alice_doc["id"])
    assert service.search("alice", "astronomy") == []
    assert len(service.list_documents("bob")) == 1


def test_duplicate_upload_is_idempotent_only_within_the_owner(tmp_path):
    service = RagService(tmp_path)
    first = service.ingest("alice", "note.txt", b"knowledge")

    assert service.ingest("alice", "note.txt", b"knowledge") == first
    assert len(service.list_documents("alice")) == 1
    assert service.ingest("bob", "note.txt", b"knowledge")["id"] != first["id"]
    assert service.ingest("alice", "note.txt", b"updated knowledge")["id"] != first["id"]


def test_untrusted_filename_is_display_name_and_never_a_storage_path(tmp_path):
    service = RagService(tmp_path / "data")
    document = service.ingest("alice", r"..\..\private\note.md", b"knowledge")

    assert document["name"] == "note.md"
    assert not (tmp_path / "private").exists()
    assert not (tmp_path / "data" / "note.md").exists()


@pytest.mark.parametrize(
    ("name", "content"),
    [
        ("script.exe", b"knowledge"),
        ("empty.txt", b" \n\t"),
        ("binary.txt", b"a\x00b\x01c"),
        ("bad.txt", b"\xff\xff"),
        ("replacement.txt", "文本\ufffd".encode()),
        ("bad.pdf", b"this is not a PDF"),
        ("bad.docx", b"this is not a document"),
        ("bad\x00.txt", b"knowledge"),
        ("large.txt", b"x" * (5 * 1024 * 1024 + 1)),
    ],
    ids=["extension", "empty", "binary", "encoding", "replacement", "pdf", "docx", "name", "size"],
)
def test_invalid_uploads_are_rejected_without_persisting(tmp_path, name, content):
    service = RagService(tmp_path)

    with pytest.raises(ValueError):
        service.ingest("alice", name, content)

    assert service.list_documents("alice") == []


def test_gb18030_and_utf16_text_preserve_chinese(tmp_path):
    service = RagService(tmp_path)
    service.ingest("alice", "old.txt", "中文知识资料".encode("gb18030"))
    service.ingest("alice", "unicode.txt", "中文知识资料".encode("utf-16"))

    assert len(service.search("alice", "知识")) == 2


def test_long_text_is_chunked_without_losing_trailing_content(tmp_path):
    service = RagService(tmp_path)
    text = ("这是中间资料。" * 200) + "结尾包含木星卫星。"
    document = service.ingest("alice", "long.md", text.encode())

    assert document["chunks"] > 1
    assert "木星卫星" in service.search("alice", "木星卫星", top_k=1)[0]["text"]
    assert all(len(result["text"]) <= 512 for result in service.search("alice", "资料"))


def test_docx_paragraphs_and_tables_are_searchable(tmp_path):
    document = Document()
    document.add_paragraph("苹果是一种水果。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "柑橘"
    table.cell(0, 1).text = "维生素"
    stream = BytesIO()
    document.save(stream)
    service = RagService(tmp_path)

    service.ingest("alice", "food.docx", stream.getvalue())

    assert "苹果" in service.search("alice", "苹果")[0]["text"]
    assert "维生素" in service.search("alice", "柑橘")[0]["text"]


@pytest.mark.parametrize("owner", ["", "   ", "x" * 257])
def test_owner_identifier_is_validated(tmp_path, owner):
    service = RagService(tmp_path)

    with pytest.raises(ValueError):
        service.ingest(owner, "note.txt", b"knowledge")
    with pytest.raises(ValueError):
        service.list_documents(owner)


def test_queries_validate_length_and_result_limit(tmp_path):
    service = RagService(tmp_path)
    assert service.search("alice", "  ") == []

    for top_k in (0, -1, 21, True):
        with pytest.raises(ValueError):
            service.search("alice", "knowledge", top_k=top_k)
    with pytest.raises(ValueError):
        service.search("alice", "x" * 16001)


def test_semantic_mode_does_not_download_missing_model(tmp_path):
    with pytest.raises(ValueError, match="本地"):
        RagService(tmp_path, embedding_model="BAAI/bge-small-zh-v1.5")


def pdf_bytes(text=None, encrypted=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 200 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = stream
    if encrypted:
        writer.encrypt("password")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_text_is_parsed_and_retrieved(tmp_path):
    service = RagService(tmp_path)

    document = service.ingest("alice", "notes.pdf", pdf_bytes("Astronomy studies Jupiter."))

    assert service.search("alice", "Jupiter")[0]["document_id"] == document["id"]


@pytest.mark.parametrize(
    "content", [pdf_bytes(), pdf_bytes("secret", encrypted=True)], ids=["blank", "encrypted"]
)
def test_blank_or_encrypted_pdf_has_explicit_validation_error(tmp_path, content):
    service = RagService(tmp_path)

    with pytest.raises(ValueError, match="OCR|加密"):
        service.ingest("alice", "notes.pdf", content)

    assert service.list_documents("alice") == []


def test_concurrent_duplicate_uploads_are_atomic(tmp_path):
    service = RagService(tmp_path)

    with ThreadPoolExecutor(max_workers=4) as executor:
        documents = list(
            executor.map(lambda _: service.ingest("alice", "note.txt", b"knowledge"), range(8))
        )

    assert len({document["id"] for document in documents}) == 1
    assert len(service.list_documents("alice")) == 1


def test_single_chinese_character_can_find_a_matching_document(tmp_path):
    service = RagService(tmp_path)
    service.ingest("alice", "pets.txt", "猫喜欢休息。".encode())

    assert "猫" in service.search("alice", "猫")[0]["text"]


def test_small_compressed_pdf_cannot_expand_into_unbounded_content(tmp_path):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    stream = DecodedStreamObject()
    stream.set_data(b" " * (4 * 1024 * 1024 + 1))
    page[NameObject("/Contents")] = stream.flate_encode()
    output = BytesIO()
    writer.write(output)
    assert len(output.getvalue()) < 10000
    service = RagService(tmp_path)
    with pytest.raises(ValueError, match="解压|内容流"):
        service.ingest("alice", "compressed.pdf", output.getvalue())
    assert service.list_documents("alice") == []
