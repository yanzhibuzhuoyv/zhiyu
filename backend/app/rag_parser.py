"""Bounded document parsing; uploaded names are never used as file paths."""

from io import BytesIO
from pathlib import PurePosixPath
from unicodedata import category
from zipfile import BadZipFile, ZipFile

from docx import Document
from lxml.etree import XMLSyntaxError
from pypdf import PdfReader, apply_configuration
from pypdf.errors import LimitReachedError, PyPdfError

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 1_000_000
MAX_PDF_PAGES = 500
MAX_DOCX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_PDF_STREAM_BYTES = 4 * 1024 * 1024
MAX_PDF_CONTENT_BYTES = 16 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


def safe_document_name(name: str) -> str:
    if not isinstance(name, str) or not name or len(name) > 512:
        raise ValueError("文件名为空或过长")
    if any(category(char).startswith("C") for char in name):
        raise ValueError("文件名包含非法字符")
    cleaned = PurePosixPath(name.replace("\\", "/")).name.strip()
    if not cleaned or cleaned in {".", ".."} or len(cleaned) > 200:
        raise ValueError("文件名无效或过长")
    if PurePosixPath(cleaned).suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("仅支持 TXT、Markdown、PDF 和 DOCX 文件")
    return cleaned


def extract_text(name: str, content: bytes) -> str:
    if not isinstance(content, bytes) or not content:
        raise ValueError("文件内容为空")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("文件不得超过 5 MB")
    extension = PurePosixPath(name).suffix.lower()
    parsers = {".txt": _decode_text, ".md": _decode_text, ".pdf": _parse_pdf, ".docx": _parse_docx}
    if extension not in parsers:
        raise ValueError("不支持的文件格式")
    text = parsers[extension](content).replace("\r\n", "\n").replace("\r", "\n").strip()
    _validate_text(text)
    return text


def _validate_text(text: str) -> None:
    if not text:
        raise ValueError("未提取到文字；扫描版 PDF 需要先进行 OCR")
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError("文档提取后的文字超过 100 万字，请拆分后上传")
    if "\ufffd" in text or any(category(char) == "Cc" and char not in "\n\t" for char in text):
        raise ValueError("文档包含二进制字符或乱码，请转换为 UTF-8 文字后上传")


def _decode_text(content: bytes) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ("utf-16",)
    else:
        encodings = ("utf-8-sig", "gb18030")
    for encoding in encodings:
        try:
            return content.decode(encoding, errors="strict")
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本编码，请保存为 UTF-8 后上传")


def _parse_pdf(content: bytes) -> str:
    try:
        with apply_configuration(
            maximum_declared_stream_length=MAX_UPLOAD_BYTES,
            array_based_stream_maximum_output_length=MAX_PDF_STREAM_BYTES,
            zlib_maximum_output_length=MAX_PDF_STREAM_BYTES,
            lzw_maximum_output_length=MAX_PDF_STREAM_BYTES,
            run_length_maximum_output_length=MAX_PDF_STREAM_BYTES,
            image_maximum_buffer_size=MAX_PDF_STREAM_BYTES,
            xform_maximum_invocations_per_extraction=256,
            page_tree_maximum_entries=2000,
            page_tree_maximum_depth=32,
        ):
            return _read_pdf_text(content)
    except LimitReachedError as exc:
        raise ValueError("PDF 解压后的内容流过大或结构过于复杂，请拆分后上传") from exc
    except (PyPdfError, OSError, KeyError, IndexError, TypeError) as exc:
        raise ValueError("无法解析 PDF，请检查文件是否损坏") from exc


def _read_pdf_text(content: bytes) -> str:
    reader = PdfReader(BytesIO(content), strict=True)
    if reader.is_encrypted:
        raise ValueError("不支持加密 PDF，请先解密")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError("PDF 不得超过 500 页")
    parts: list[str] = []
    length, stream_bytes = 0, 0
    for page in reader.pages:
        stream = page.get_contents()
        stream_bytes += len(stream.get_data()) if stream else 0
        if stream_bytes > MAX_PDF_CONTENT_BYTES:
            raise ValueError("PDF 解压后的内容流总量过大，请拆分后上传")
        text = page.extract_text() or ""
        length += len(text)
        if length > MAX_TEXT_CHARS:
            raise ValueError("PDF 提取后的文字过长，请拆分后上传")
        parts.append(text)
    return "\n".join(parts)


def _parse_docx(content: bytes) -> str:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 1024
                or sum(item.file_size for item in entries) > MAX_DOCX_EXPANDED_BYTES
            ):
                raise ValueError("DOCX 解压后的内容过大，请拆分后上传")
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join(paragraphs + cells)
    except (BadZipFile, KeyError, OSError, XMLSyntaxError, TypeError) as exc:
        raise ValueError("无法解析 DOCX，请检查文件是否损坏") from exc
